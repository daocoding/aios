#!/usr/bin/env python3
"""Build Excel templates from apps/content-review/static/tools.json (the same spec the web app uses).
Writes templates/xlsx/<tool>.xlsx (one per tool) and templates/aios-operating-sheets.xlsx (all tools, one tab each).
Formulas are real Excel formulas; select columns get dropdown validation; example rows are included.
Requires: pip install openpyxl"""
import json, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
spec = json.load(open(os.path.join(ROOT, "apps/content-review/static/tools.json")))
FMT = {"pct": "0.0%", "money": '"$"#,##0.00', "num2": "0.00", "int": "#,##0", "number": "General"}
HEAD = PatternFill("solid", fgColor="111316"); FHEAD = PatternFill("solid", fgColor="0F766E"); FCELL = PatternFill("solid", fgColor="F2FAF8")
ROWS = 200

def tab_name(t):
    """Same rule as sheets.js S.tabName: explicit 'tab', else title without parentheticals; max 31 chars."""
    n = t.get("tab") or re.sub(r"\s*\(.*?\)\s*", " ", t["title"]).strip()
    return re.sub(r"[\\/?*\[\]:]", "-", n)[:31]

def fill(ws, t):
    cols = t["columns"]; idx = {c["key"]: i for i, c in enumerate(cols)}
    ws.append([c["label"] + (" ƒ" if c["type"] == "formula" else "") for c in cols])
    for i, c in enumerate(cols, 1):
        h = ws.cell(1, i); h.font = Font(bold=True, color="FFFFFF"); h.fill = FHEAD if c["type"] == "formula" else HEAD
        h.alignment = Alignment(wrap_text=True, vertical="top"); ws.column_dimensions[L(i)].width = max(12, (c.get("w", 110)) / 7)
    ws.row_dimensions[1].height = 48; ws.freeze_panes = "B2"
    samples = t.get("sample", [])
    for r in range(2, ROWS + 2):
        s = samples[r - 2] if r - 2 < len(samples) else {}
        for i, c in enumerate(cols, 1):
            cell = ws.cell(r, i)
            if c["type"] == "formula":
                f = re.sub(r"\{(\w+)\}", lambda m: f"{L(idx[m.group(1)] + 1)}{r}", c["f"])
                first = L(idx[next(k["key"] for k in cols if k["type"] != "formula")] + 1)
                cell.value = f'=IF({first}{r}="","",{f})'; cell.fill = FCELL
                cell.number_format = FMT.get(c.get("fmt", ""), "General")
            else:
                v = s.get(c["key"]);
                if v is not None: cell.value = v
                cell.number_format = FMT.get(c["type"], "General")
    for i, c in enumerate(cols, 1):
        if c["type"] == "select":
            dv = DataValidation(type="list", formula1='"' + ",".join(o for o in c["options"] if o) + '"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f"{L(i)}2:{L(i)}{ROWS + 1}")
    tot = t.get("totals")
    if tot:
        r = ROWS + 3; ws.cell(r, 1, "TOTAL").font = Font(bold=True)
        for k, how in tot.items():
            col = L(idx[k] + 1); c = ws.cell(r, idx[k] + 1, f"={'COUNTA' if how == 'COUNT' else 'SUM'}({col}2:{col}{ROWS + 1})")
            c.font = Font(bold=True); c.number_format = FMT.get(cols[idx[k]].get("fmt", cols[idx[k]]["type"]), "General")
    for j, k in enumerate(t.get("kpis", [])):
        r = ROWS + 5 + j
        f = re.sub(r"(SUM|COUNT|AVG)\((\w+)\)", lambda m: f"{ {'SUM':'SUM','COUNT':'COUNTA','AVG':'AVERAGE'}[m.group(1)]}({L(idx[m.group(2)] + 1)}2:{L(idx[m.group(2)] + 1)}{ROWS + 1})", k["f"])
        ws.cell(r, 1, k["label"]).font = Font(bold=True)
        c = ws.cell(r, 2, f"=IFERROR({f},0)"); c.number_format = FMT.get(k.get("fmt"), "General")

def about(ws, t):
    ws.append(["aiOS Grok Edition · " + t["title"]]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([t["purpose"]]); ws.append([f"Module {t['module']} · see modules/ in https://github.com/daocoding/aios"])
    ws.append(["Columns marked ƒ are formulas. Rows 2-4 hold examples; delete them before real use."]); ws.column_dimensions["A"].width = 120

out = os.path.join(ROOT, "templates", "xlsx"); os.makedirs(out, exist_ok=True)
book = Workbook(); book.remove(book.active)
for t in spec["tools"]:
    wb = Workbook(); ws = wb.active; ws.title = tab_name(t); fill(ws, t); about(wb.create_sheet("About"), t)
    wb.save(os.path.join(out, t["id"] + ".xlsx"))
    fill(book.create_sheet(tab_name(t)), t)
book.save(os.path.join(ROOT, "templates", "aios-operating-sheets.xlsx"))
print("built", len(spec["tools"]), "templates")
